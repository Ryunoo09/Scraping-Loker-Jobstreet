"""Writer Excel lokal berbasis ``openpyxl`` (append-only, simpan atomik)."""

from __future__ import annotations

import logging
import os
import tempfile
from collections.abc import Sequence
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from ..exceptions import WriterError
from ..models import JobLinkedIn
from .base import HEADERS, job_to_row

logger = logging.getLogger(__name__)

_SHEET_TITLE = "LinkedIn"
_LINK_COLUMN = 5  # kolom E: "Link Lowongan" (1-based)


class ExcelWriter:
    """Tulis hasil ke satu file ``.xlsx`` lokal secara append-only."""

    def __init__(self, path: Path) -> None:
        self._path = Path(path)
        self._header_ready = False

    def ensure_header(self) -> None:
        """Buat file + header bila belum ada (idempoten)."""
        if self._header_ready:
            return
        workbook, worksheet = self._load()
        if _is_empty(worksheet):
            _write_row(worksheet, 1, list(HEADERS))
            self._save(workbook)
            logger.info("Membuat file Excel + header: %s", self._path)
        self._header_ready = True

    def append(self, jobs: Sequence[JobLinkedIn]) -> None:
        """Tambahkan ``jobs`` sebagai baris baru di bawah data yang ada."""
        if not jobs:
            return
        workbook, worksheet = self._load()
        if not self._header_ready and _is_empty(worksheet):
            _write_row(worksheet, 1, list(HEADERS))
        self._header_ready = True

        row_index = _next_row(worksheet)
        for job in jobs:
            _write_row(worksheet, row_index, job_to_row(job))
            _apply_hyperlink(worksheet, row_index, job.url)
            row_index += 1

        self._save(workbook)
        logger.info("Menulis %d baris ke %s", len(jobs), self._path)

    def _load(self) -> tuple[Workbook, Worksheet]:
        """Muat workbook yang ada, atau siapkan workbook baru."""
        if self._path.is_file():
            try:
                workbook = load_workbook(self._path)
            except (OSError, BadZipFile) as exc:
                raise WriterError(f"File Excel tidak bisa dibaca ({self._path}): {exc}") from exc
            return workbook, workbook.active

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = _SHEET_TITLE
        return workbook, worksheet

    def _save(self, workbook: Workbook) -> None:
        """Simpan secara atomik; terjemahkan kegagalan (mis. file terkunci)."""
        self._path.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp_name = tempfile.mkstemp(
            dir=self._path.parent, prefix=f".{self._path.stem}_", suffix=".tmp"
        )
        os.close(fd)
        tmp_path = Path(tmp_name)
        try:
            workbook.save(tmp_path)
            os.replace(tmp_path, self._path)
        except OSError as exc:
            raise WriterError(
                f"Gagal menyimpan {self._path} — pastikan file tidak sedang dibuka di Excel: {exc}"
            ) from exc
        finally:
            tmp_path.unlink(missing_ok=True)


def _is_empty(worksheet: Worksheet) -> bool:
    """``True`` bila worksheet belum punya isi apa pun."""
    return worksheet.max_row == 1 and worksheet.cell(row=1, column=1).value is None


def _next_row(worksheet: Worksheet) -> int:
    """Nomor baris berikutnya untuk data; header selalu menempati baris 1."""
    return int(worksheet.max_row) + 1


def _write_row(worksheet: Worksheet, row_index: int, values: Sequence[str | int]) -> None:
    """Tulis ``values`` mulai kolom A pada ``row_index``."""
    for column, value in enumerate(values, start=1):
        worksheet.cell(row=row_index, column=column, value=value)


def _apply_hyperlink(worksheet: Worksheet, row_index: int, url: str) -> None:
    """Jadikan sel "Link Lowongan" bisa diklik."""
    if not url:
        return
    cell = worksheet.cell(row=row_index, column=_LINK_COLUMN)
    cell.hyperlink = url
    cell.style = "Hyperlink"
