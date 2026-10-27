"""Test CLI LinkedIn: flag, ringkasan, dan orkestrasi per halaman (tanpa jaringan)."""

from __future__ import annotations

from collections import Counter
from collections.abc import Iterator
from pathlib import Path

import pytest
from openpyxl import load_workbook

import linkedin_scraper.cli as cli
from linkedin_scraper.cli import _describe_output, _process_cards, build_parser, main
from linkedin_scraper.config import (
    AppConfig,
    ExcelConfig,
    FiltersConfig,
    OutputConfig,
    ScoringConfig,
    SearchConfig,
)
from linkedin_scraper.exceptions import RateLimitError, ScraperError
from linkedin_scraper.models import JobLinkedIn
from linkedin_scraper.parsers import parse_search_page
from linkedin_scraper.storage import SeenJobsDB

FIXTURES = Path(__file__).parent / "fixtures" / "linkedin"
CARDS_HTML = (FIXTURES / "guest_cards_01.html").read_text(encoding="utf-8")
DETAIL_HTML = (FIXTURES / "detail_01.html").read_text(encoding="utf-8")
EXAMPLE = Path(__file__).resolve().parent.parent / "config_linkedin.example.yaml"


class _FakeSource:
    """Pengganti HttpSource: kembalikan fixture, tanpa jaringan."""

    def __init__(self, *, fail_detail: bool = False, fail_search: bool = False) -> None:
        self._fail_detail = fail_detail
        self._fail_search = fail_search
        self.detail_calls = 0

    def __enter__(self) -> _FakeSource:
        return self

    def __exit__(self, *exc_info: object) -> None:
        return None

    def search_page_html(self, keyword: str, page: int) -> str:
        if self._fail_search:
            raise RateLimitError("429")
        return CARDS_HTML

    def fetch_detail_html(self, job_id: str) -> str:
        self.detail_calls += 1
        if self._fail_detail:
            raise ScraperError("network error")
        return DETAIL_HTML


def _cfg(*, fetch_detail: bool = True, min_score: int = 0) -> AppConfig:
    return AppConfig(
        search=SearchConfig(keywords=["Developer"], fetch_detail=fetch_detail),
        filters=FiltersConfig(),
        scoring=ScoringConfig(min_score=min_score),
        output=OutputConfig(type="excel", excel=ExcelConfig()),
    )


@pytest.fixture
def cards() -> list[JobLinkedIn]:
    return parse_search_page(CARDS_HTML)[:2]


@pytest.fixture
def db() -> Iterator[SeenJobsDB]:
    with SeenJobsDB(":memory:") as database:
        yield database


# --------------------------------------------------------------------------- #
# Flag parsing
# --------------------------------------------------------------------------- #
def test_build_parser_defaults() -> None:
    args = build_parser().parse_args([])
    assert args.dry_run is False
    assert args.limit == 0
    assert args.reset_rejected is False


def test_build_parser_flags() -> None:
    args = build_parser().parse_args(["--dry-run", "--limit", "5", "--reset-rejected"])
    assert args.dry_run is True
    assert args.limit == 5
    assert args.reset_rejected is True


def test_describe_output_excel() -> None:
    cfg = _cfg()
    assert _describe_output(cfg) == "data/linkedin_jobs.xlsx"


# --------------------------------------------------------------------------- #
# _process_cards
# --------------------------------------------------------------------------- #
def test_process_cards_passes_and_marks(cards: list[JobLinkedIn], db: SeenJobsDB) -> None:
    stats: Counter[str] = Counter()
    matches: list[JobLinkedIn] = []
    source = _FakeSource()

    _process_cards(cards, _cfg(), db, source, stats, matches, limit=0, dry_run=False)

    assert len(matches) == 2
    assert stats["lolos"] == 2
    assert stats["dilihat"] == 2
    assert db.get_status(cards[0].job_id) == "matched"
    assert matches[0].detail_complete is True
    assert matches[0].description is not None


def test_process_cards_skips_duplicates(cards: list[JobLinkedIn], db: SeenJobsDB) -> None:
    stats: Counter[str] = Counter()
    matches: list[JobLinkedIn] = []
    source = _FakeSource()

    _process_cards(cards, _cfg(), db, source, stats, matches, limit=0, dry_run=False)
    _process_cards(cards, _cfg(), db, source, stats, matches, limit=0, dry_run=False)

    assert stats["duplikat"] == 2
    assert len(matches) == 2


def test_process_cards_continues_when_detail_fails(
    cards: list[JobLinkedIn], db: SeenJobsDB
) -> None:
    stats: Counter[str] = Counter()
    matches: list[JobLinkedIn] = []
    source = _FakeSource(fail_detail=True)

    _process_cards(cards, _cfg(), db, source, stats, matches, limit=0, dry_run=False)

    assert len(matches) == 2
    assert matches[0].description is None
    assert matches[0].detail_complete is False


def test_process_cards_rejects_low_score(cards: list[JobLinkedIn], db: SeenJobsDB) -> None:
    stats: Counter[str] = Counter()
    matches: list[JobLinkedIn] = []
    source = _FakeSource()

    _process_cards(cards, _cfg(min_score=99), db, source, stats, matches, limit=0, dry_run=False)

    assert matches == []
    assert stats["ditolak"] == 2
    assert db.get_status(cards[0].job_id) == "rejected_low_score"


def test_process_cards_respects_limit(cards: list[JobLinkedIn], db: SeenJobsDB) -> None:
    stats: Counter[str] = Counter()
    matches: list[JobLinkedIn] = []

    _process_cards(cards, _cfg(), db, _FakeSource(), stats, matches, limit=1, dry_run=False)

    assert stats["dilihat"] == 1


def test_process_cards_dry_run_does_not_mark(cards: list[JobLinkedIn], db: SeenJobsDB) -> None:
    stats: Counter[str] = Counter()
    matches: list[JobLinkedIn] = []

    _process_cards(cards, _cfg(), db, _FakeSource(), stats, matches, limit=0, dry_run=True)

    assert len(matches) == 2
    assert db.has(cards[0].job_id) is False


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #
def test_main_check_config(capsys: pytest.CaptureFixture[str]) -> None:
    exit_code = main(["--config", str(EXAMPLE), "--check-config"])
    assert exit_code == 0
    assert "Config valid" in capsys.readouterr().out


def test_main_invalid_config(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    bad = tmp_path / "bad.yaml"
    bad.write_text("search: {}\n", encoding="utf-8")
    assert main(["--config", str(bad), "--check-config"]) == 2
    assert "Config tidak valid" in capsys.readouterr().err


# --------------------------------------------------------------------------- #
# run() end-to-end (HttpSource di-fake, filesystem diarahkan ke tmp)
# --------------------------------------------------------------------------- #
def _write_config(tmp_path: Path) -> Path:
    out = (tmp_path / "out.xlsx").as_posix()
    cfg_path = tmp_path / "config.yaml"
    cfg_path.write_text(
        "search:\n"
        '  keywords: ["Developer"]\n'
        '  location: "Indonesia"\n'
        "  max_pages_per_query: 1\n"
        "output:\n"
        '  type: "excel"\n'
        "  excel:\n"
        f'    file_path: "{out}"\n'
        "scoring:\n"
        "  min_score: 0\n",
        encoding="utf-8",
    )
    return cfg_path


def _patch_paths(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    monkeypatch.setattr(cli, "_DATA", tmp_path / "data")
    monkeypatch.setattr(cli, "_LOG_DIR", tmp_path / "logs")


def test_run_writes_excel(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg_path = _write_config(tmp_path)
    _patch_paths(monkeypatch, tmp_path)
    monkeypatch.setattr(cli, "HttpSource", lambda *a, **k: _FakeSource())

    args = build_parser().parse_args(["--config", str(cfg_path), "--limit", "2"])
    assert cli.run(cfg_path, args) == 0

    worksheet = load_workbook(tmp_path / "out.xlsx").active
    assert worksheet.max_row == 3  # header + 2 baris
    assert (tmp_path / "data" / "linkedin_seen.db").is_file()


def test_run_dry_run_writes_nothing(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg_path = _write_config(tmp_path)
    _patch_paths(monkeypatch, tmp_path)
    monkeypatch.setattr(cli, "HttpSource", lambda *a, **k: _FakeSource())

    args = build_parser().parse_args(["--config", str(cfg_path), "--dry-run", "--limit", "2"])
    assert cli.run(cfg_path, args) == 0
    assert (tmp_path / "out.xlsx").is_file() is False


def test_run_stops_on_rate_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    cfg_path = _write_config(tmp_path)
    _patch_paths(monkeypatch, tmp_path)
    monkeypatch.setattr(cli, "HttpSource", lambda *a, **k: _FakeSource(fail_search=True))

    args = build_parser().parse_args(["--config", str(cfg_path)])
    assert cli.run(cfg_path, args) == 1
