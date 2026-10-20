"""Test untuk penulis output (Excel, Google Sheets, null) — M4.

Google Sheets di-mock: test tidak menyentuh jaringan (spec §15).
"""

from __future__ import annotations

import logging
from collections.abc import Callable
from pathlib import Path

import gspread
import pytest
from gspread.exceptions import GSpreadException
from gspread.utils import ValueInputOption
from openpyxl import load_workbook

from job_scraper.config import ExcelConfig, GoogleSheetsConfig, OutputConfig
from job_scraper.exceptions import WriterError
from job_scraper.models import Job
from job_scraper.writers import HEADERS, ExcelWriter, NullWriter, build_writer, job_to_row
from job_scraper.writers.sheets import GoogleSheetsWriter

# --------------------------------------------------------------------------- #
# Helper & fake gspread
# --------------------------------------------------------------------------- #


class FakeWorksheet:
    """Pengganti ``gspread.Worksheet`` untuk test offline."""

    def __init__(self, first_row: list[str] | None = None) -> None:
        self.first_row = list(first_row) if first_row else None
        self.inserted: list[list[str]] = []
        self.appended: list[tuple[list[list[str | int]], ValueInputOption | None]] = []
        self.row_values_calls = 0
        self.row_values_error: Exception | None = None
        self.append_error: Exception | None = None

    def row_values(self, row: int) -> list[str]:
        self.row_values_calls += 1
        if self.row_values_error is not None:
            raise self.row_values_error
        return list(self.first_row or [])

    def insert_row(self, values: list[str], index: int = 1) -> None:
        self.inserted.append(list(values))
        self.first_row = list(values)

    def append_rows(
        self,
        values: list[list[str | int]],
        value_input_option: ValueInputOption | None = None,
    ) -> None:
        if self.append_error is not None:
            raise self.append_error
        self.appended.append(([list(row) for row in values], value_input_option))


def _patch_service_account(monkeypatch: pytest.MonkeyPatch, worksheet: FakeWorksheet) -> None:
    """Ganti ``gspread.service_account`` dengan klien palsu."""

    class _FakeSpreadsheet:
        def worksheet(self, name: str) -> FakeWorksheet:
            return worksheet

    class _FakeClient:
        def open_by_key(self, key: str) -> _FakeSpreadsheet:
            return _FakeSpreadsheet()

    def _service_account(*, filename: str, **kwargs: object) -> _FakeClient:
        return _FakeClient()

    monkeypatch.setattr(gspread, "service_account", _service_account)


def _gs_config() -> GoogleSheetsConfig:
    return GoogleSheetsConfig(
        spreadsheet_id="spreadsheet-123",
        worksheet_name="Lowongan",
        credentials_file="credentials/service_account.json",
    )


def _raise_permission(*args: object, **kwargs: object) -> None:
    raise PermissionError("file dipakai proses lain")


# --------------------------------------------------------------------------- #
# job_to_row
# --------------------------------------------------------------------------- #


def test_job_to_row_maps_all_columns(make_job: Callable[..., Job]) -> None:
    job = make_job(
        job_id="42",
        title="Junior Web Developer",
        company="PT Contoh",
        location="Semarang",
        url="https://id.jobstreet.com/job/42",
        salary_text="Rp 5.000.000",
        work_type="Full time",
        posted_date="2026-10-01",
        score=77,
    )

    row = job_to_row(job)

    assert len(row) == len(HEADERS) == 14
    assert row[0] == job.found_date
    assert row[1] == "Junior Web Developer"
    assert row[2] == "PT Contoh"
    assert row[3] == "Semarang"
    assert row[4] == "https://id.jobstreet.com/job/42"
    assert row[5] == "Rp 5.000.000"
    assert row[6] == "Full time"
    assert row[7] == "2026-10-01"
    assert row[8] == 77
    assert row[9] == "Belum"
    assert row[10] == ""
    assert row[11] == ""  # description
    assert row[12] == ""  # requirements
    assert row[13] == "42"


def test_job_to_row_uses_empty_strings_for_missing_optionals(
    make_job: Callable[..., Job],
) -> None:
    row = job_to_row(make_job())

    assert row[5] == ""
    assert row[6] == ""
    assert row[7] == ""


def test_job_to_row_formats_posted_at_as_wib(make_job: Callable[..., Job]) -> None:
    job = make_job(posted_at="2026-10-13T02:58:53.000Z", posted_date="12 jam yang lalu")

    assert job_to_row(job)[7] == "2026-10-13 09:58"


def test_job_to_row_falls_back_to_relative_posted_date(make_job: Callable[..., Job]) -> None:
    job = make_job(posted_date="12 jam yang lalu")

    assert job_to_row(job)[7] == "12 jam yang lalu"


# --------------------------------------------------------------------------- #
# ExcelWriter
# --------------------------------------------------------------------------- #


def test_excel_ensure_header_creates_file_with_header(tmp_path: Path) -> None:
    path = tmp_path / "jobs.xlsx"

    ExcelWriter(path).ensure_header()

    assert path.is_file()
    assert [cell.value for cell in load_workbook(path).active[1]] == list(HEADERS)


def test_excel_ensure_header_is_idempotent(tmp_path: Path) -> None:
    path = tmp_path / "jobs.xlsx"

    writer = ExcelWriter(path)
    writer.ensure_header()
    writer.ensure_header()  # cabang early-return (header sudah siap)
    ExcelWriter(path).ensure_header()  # tetap idempoten lintas instance

    assert load_workbook(path).active.max_row == 1


def test_excel_append_adds_rows_below_header(tmp_path: Path, make_job: Callable[..., Job]) -> None:
    path = tmp_path / "jobs.xlsx"
    writer = ExcelWriter(path)
    writer.ensure_header()

    writer.append([make_job(job_id="1", title="A"), make_job(job_id="2", title="B")])

    worksheet = load_workbook(path).active
    assert [cell.value for cell in worksheet[1]] == list(HEADERS)
    assert worksheet.max_row == 3
    assert worksheet.cell(row=2, column=2).value == "A"
    assert worksheet.cell(row=3, column=2).value == "B"


def test_excel_append_without_ensure_header_still_writes_header(
    tmp_path: Path, make_job: Callable[..., Job]
) -> None:
    path = tmp_path / "jobs.xlsx"

    ExcelWriter(path).append([make_job()])

    worksheet = load_workbook(path).active
    assert [cell.value for cell in worksheet[1]] == list(HEADERS)
    assert worksheet.max_row == 2


def test_excel_append_empty_sequence_does_nothing(tmp_path: Path) -> None:
    path = tmp_path / "jobs.xlsx"

    ExcelWriter(path).append([])

    assert not path.exists()


def test_excel_append_is_append_only(tmp_path: Path, make_job: Callable[..., Job]) -> None:
    path = tmp_path / "jobs.xlsx"
    ExcelWriter(path).append([make_job(job_id="1", title="Pertama")])

    # Simulasi pengguna mengedit kolom status lamaran secara manual.
    workbook = load_workbook(path)
    worksheet = workbook.active
    worksheet.cell(row=2, column=10).value = "Sudah Dikirim"
    workbook.save(path)

    ExcelWriter(path).append([make_job(job_id="2", title="Kedua")])

    worksheet = load_workbook(path).active
    assert worksheet.max_row == 3
    assert worksheet.cell(row=2, column=2).value == "Pertama"
    assert worksheet.cell(row=2, column=10).value == "Sudah Dikirim"
    assert worksheet.cell(row=3, column=2).value == "Kedua"


def test_excel_sets_clickable_hyperlink(tmp_path: Path, make_job: Callable[..., Job]) -> None:
    path = tmp_path / "jobs.xlsx"
    url = "https://id.jobstreet.com/job/999"

    ExcelWriter(path).append([make_job(url=url)])

    cell = load_workbook(path).active.cell(row=2, column=5)
    assert cell.value == url
    assert cell.hyperlink is not None
    assert cell.hyperlink.target == url


def test_excel_creates_parent_directory(tmp_path: Path, make_job: Callable[..., Job]) -> None:
    path = tmp_path / "nested" / "dir" / "jobs.xlsx"

    ExcelWriter(path).append([make_job()])

    assert path.is_file()


def test_excel_job_without_url_skips_hyperlink(
    tmp_path: Path, make_job: Callable[..., Job]
) -> None:
    path = tmp_path / "jobs.xlsx"

    ExcelWriter(path).append([make_job(url="")])

    cell = load_workbook(path).active.cell(row=2, column=5)
    assert not cell.value  # string kosong tersimpan sebagai sel kosong
    assert cell.hyperlink is None


def test_excel_locked_file_raises_writer_error(
    tmp_path: Path, make_job: Callable[..., Job], monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "jobs.xlsx"
    monkeypatch.setattr("job_scraper.writers.excel.os.replace", _raise_permission)

    with pytest.raises(WriterError, match="tidak sedang dibuka"):
        ExcelWriter(path).append([make_job()])

    assert list(tmp_path.glob(".*tmp")) == []


def test_excel_unreadable_file_raises_writer_error(tmp_path: Path) -> None:
    path = tmp_path / "jobs.xlsx"
    path.write_text("bukan file xlsx", encoding="utf-8")

    with pytest.raises(WriterError, match="tidak bisa dibaca"):
        ExcelWriter(path).ensure_header()


# --------------------------------------------------------------------------- #
# NullWriter
# --------------------------------------------------------------------------- #


def test_null_writer_is_noop(make_job: Callable[..., Job]) -> None:
    writer = NullWriter()

    writer.ensure_header()
    writer.append([make_job()])


# --------------------------------------------------------------------------- #
# build_writer
# --------------------------------------------------------------------------- #


def test_build_writer_dry_run_returns_null_writer(
    tmp_path: Path, make_job: Callable[..., Job]
) -> None:
    output = OutputConfig(type="excel", excel=ExcelConfig(file_path=str(tmp_path / "jobs.xlsx")))

    writer = build_writer(output, dry_run=True)
    writer.append([make_job()])

    assert isinstance(writer, NullWriter)
    assert not (tmp_path / "jobs.xlsx").exists()


def test_build_writer_excel_prepares_header(tmp_path: Path) -> None:
    path = tmp_path / "jobs.xlsx"
    output = OutputConfig(type="excel", excel=ExcelConfig(file_path=str(path)))

    writer = build_writer(output)

    assert isinstance(writer, ExcelWriter)
    assert [cell.value for cell in load_workbook(path).active[1]] == list(HEADERS)


def test_build_writer_google_sheets_prepares_header(monkeypatch: pytest.MonkeyPatch) -> None:
    worksheet = FakeWorksheet()
    _patch_service_account(monkeypatch, worksheet)
    output = OutputConfig(type="google_sheets", google_sheets=_gs_config())

    writer = build_writer(output)

    assert isinstance(writer, GoogleSheetsWriter)
    assert worksheet.inserted == [list(HEADERS)]


def test_build_writer_google_sheets_without_config_raises() -> None:
    output = OutputConfig.model_construct(type="google_sheets")

    with pytest.raises(WriterError, match="google_sheets"):
        build_writer(output)


def test_build_writer_excel_without_config_raises() -> None:
    output = OutputConfig.model_construct(type="excel")

    with pytest.raises(WriterError, match=r"output\.excel"):
        build_writer(output)


# --------------------------------------------------------------------------- #
# GoogleSheetsWriter (di-mock)
# --------------------------------------------------------------------------- #


def test_sheets_ensure_header_inserts_when_empty(monkeypatch: pytest.MonkeyPatch) -> None:
    worksheet = FakeWorksheet()
    _patch_service_account(monkeypatch, worksheet)

    GoogleSheetsWriter(_gs_config()).ensure_header()

    assert worksheet.inserted == [list(HEADERS)]


def test_sheets_ensure_header_idempotent(monkeypatch: pytest.MonkeyPatch) -> None:
    worksheet = FakeWorksheet(first_row=list(HEADERS))
    _patch_service_account(monkeypatch, worksheet)
    writer = GoogleSheetsWriter(_gs_config())

    writer.ensure_header()
    writer.ensure_header()

    assert worksheet.inserted == []
    assert worksheet.row_values_calls == 1


def test_sheets_keeps_mismatched_header(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    worksheet = FakeWorksheet(first_row=["Kolom Sendiri"])
    _patch_service_account(monkeypatch, worksheet)

    with caplog.at_level(logging.WARNING):
        GoogleSheetsWriter(_gs_config()).ensure_header()

    assert worksheet.inserted == []
    assert "tidak sesuai standar" in caplog.text


def test_sheets_append_writes_single_batch(
    monkeypatch: pytest.MonkeyPatch, make_job: Callable[..., Job]
) -> None:
    worksheet = FakeWorksheet()
    _patch_service_account(monkeypatch, worksheet)
    writer = GoogleSheetsWriter(_gs_config())

    writer.append([make_job(job_id="1"), make_job(job_id="2")])

    assert worksheet.inserted == [list(HEADERS)]
    assert len(worksheet.appended) == 1
    rows, value_input_option = worksheet.appended[0]
    assert len(rows) == 2
    assert value_input_option == ValueInputOption.user_entered
    assert rows[0][13] == "1"
    assert rows[1][13] == "2"


def test_sheets_append_empty_sequence_is_noop(monkeypatch: pytest.MonkeyPatch) -> None:
    worksheet = FakeWorksheet()
    _patch_service_account(monkeypatch, worksheet)

    GoogleSheetsWriter(_gs_config()).append([])

    assert worksheet.appended == []
    assert worksheet.inserted == []


def test_sheets_connection_error_raises_writer_error(monkeypatch: pytest.MonkeyPatch) -> None:
    def _boom(*, filename: str, **kwargs: object) -> None:
        raise GSpreadException("kredensial tidak valid")

    monkeypatch.setattr(gspread, "service_account", _boom)

    with pytest.raises(WriterError, match="Gagal terhubung"):
        GoogleSheetsWriter(_gs_config())


def test_sheets_ensure_header_error_raises_writer_error(monkeypatch: pytest.MonkeyPatch) -> None:
    worksheet = FakeWorksheet()
    worksheet.row_values_error = GSpreadException("API tidak tersedia")
    _patch_service_account(monkeypatch, worksheet)

    with pytest.raises(WriterError, match="menyiapkan header"):
        GoogleSheetsWriter(_gs_config()).ensure_header()


def test_sheets_append_error_raises_writer_error(
    monkeypatch: pytest.MonkeyPatch, make_job: Callable[..., Job]
) -> None:
    worksheet = FakeWorksheet()
    worksheet.append_error = GSpreadException("kuota habis")
    _patch_service_account(monkeypatch, worksheet)

    with pytest.raises(WriterError, match="Gagal menulis"):
        GoogleSheetsWriter(_gs_config()).append([make_job()])
