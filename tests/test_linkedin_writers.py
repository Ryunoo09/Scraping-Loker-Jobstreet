"""Test writer LinkedIn: header, pemetaan baris, dan append-only."""

from __future__ import annotations

from pathlib import Path

import gspread
import pytest
from gspread.exceptions import GSpreadException
from openpyxl import load_workbook

import linkedin_scraper.writers as writers_pkg
from linkedin_scraper.config import ExcelConfig, GoogleSheetsConfig, OutputConfig
from linkedin_scraper.exceptions import WriterError
from linkedin_scraper.models import JobLinkedIn
from linkedin_scraper.writers import HEADERS, NullWriter, build_writer
from linkedin_scraper.writers.base import job_to_row
from linkedin_scraper.writers.excel import ExcelWriter
from linkedin_scraper.writers.sheets import GoogleSheetsWriter


def _job(**overrides: object) -> JobLinkedIn:
    data: dict[str, object] = {
        "job_id": "4446754033",
        "title": "Laravel Programmer",
        "company": "PT Contoh",
        "location": "Depok",
        "url": "https://www.linkedin.com/jobs/view/4446754033",
        "posted_at": "2026-07-30",
        "description": "desc",
        "requirements": "req",
        "score": 42,
        "detail_complete": True,
    }
    data.update(overrides)
    return JobLinkedIn.model_validate(data)


def test_headers_have_fifteen_columns() -> None:
    assert len(HEADERS) == 15
    assert HEADERS[0] == "Tanggal Ditemukan"
    assert HEADERS[-1] == "Detail Lengkap?"


def test_job_to_row_maps_fields() -> None:
    row = job_to_row(_job())
    assert row[1] == "Laravel Programmer"
    assert row[7] == "2026-07-30"
    assert row[8] == 42
    assert row[9] == "Belum"
    assert row[14] == "Ya"


def test_job_to_row_falls_back_to_relative_date() -> None:
    row = job_to_row(
        _job(posted_at=None, posted_date_text="2 bulan yang lalu", detail_complete=False)
    )
    assert row[7] == "2 bulan yang lalu"
    assert row[14] == "Tidak"


def test_excel_writer_creates_header_and_appends(tmp_path: Path) -> None:
    path = tmp_path / "linkedin.xlsx"
    writer = ExcelWriter(path)
    writer.ensure_header()
    writer.append([_job()])

    worksheet = load_workbook(path).active
    assert [cell.value for cell in worksheet[1]] == list(HEADERS)
    assert worksheet.max_row == 2
    assert worksheet.cell(row=2, column=2).value == "Laravel Programmer"
    assert worksheet.cell(row=2, column=5).hyperlink.target == _job().url


def test_excel_writer_is_append_only(tmp_path: Path) -> None:
    path = tmp_path / "linkedin.xlsx"
    ExcelWriter(path).append([_job()])
    ExcelWriter(path).append([_job(job_id="2", title="Second")])

    worksheet = load_workbook(path).active
    assert worksheet.max_row == 3
    assert worksheet.cell(row=2, column=2).value == "Laravel Programmer"
    assert worksheet.cell(row=3, column=2).value == "Second"


def test_build_writer_dry_run_returns_null(tmp_path: Path) -> None:
    cfg = OutputConfig(type="excel", excel=ExcelConfig(file_path=str(tmp_path / "x.xlsx")))
    assert isinstance(build_writer(cfg, dry_run=True), NullWriter)


def test_build_writer_creates_excel_file(tmp_path: Path) -> None:
    path = tmp_path / "out.xlsx"
    cfg = OutputConfig(type="excel", excel=ExcelConfig(file_path=str(path)))
    build_writer(cfg, dry_run=False)
    assert path.is_file()


class _FakeWorksheet:
    """Worksheet palsu: cukup untuk header + append."""

    def __init__(self) -> None:
        self.header: list[str] = []
        self.rows: list[list[str | int]] = []
        self.last_value_input_option = ""

    def row_values(self, row: int) -> list[str]:
        return list(self.header)

    def insert_row(self, values: list[str], index: int) -> None:
        self.header = list(values)

    def append_rows(self, rows: list[list[str | int]], value_input_option: str = "") -> None:
        self.last_value_input_option = value_input_option
        self.rows.extend(rows)


class _FakeSpreadsheet:
    def __init__(self, worksheet: _FakeWorksheet) -> None:
        self._worksheet = worksheet

    def worksheet(self, name: str) -> _FakeWorksheet:
        return self._worksheet


class _FakeClient:
    def __init__(self, worksheet: _FakeWorksheet) -> None:
        self._worksheet = worksheet

    def open_by_key(self, key: str) -> _FakeSpreadsheet:
        return _FakeSpreadsheet(self._worksheet)


def test_sheets_writer_writes_header_and_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    worksheet = _FakeWorksheet()
    monkeypatch.setattr(gspread, "service_account", lambda filename: _FakeClient(worksheet))

    writer = GoogleSheetsWriter(GoogleSheetsConfig(spreadsheet_id="abc"))
    writer.ensure_header()
    writer.append([_job()])

    assert worksheet.header == list(HEADERS)
    assert len(worksheet.rows) == 1
    assert worksheet.rows[0][1] == "Laravel Programmer"
    assert worksheet.last_value_input_option == "RAW"


def test_sheets_writer_connection_error(monkeypatch: pytest.MonkeyPatch) -> None:
    def boom(filename: str) -> None:
        raise GSpreadException("no connection")

    monkeypatch.setattr(gspread, "service_account", boom)
    with pytest.raises(WriterError, match="Gagal terhubung"):
        GoogleSheetsWriter(GoogleSheetsConfig(spreadsheet_id="abc"))


def test_excel_writer_forces_text_for_formula_like_values(tmp_path: Path) -> None:
    path = tmp_path / "inj.xlsx"
    ExcelWriter(path).append([_job(title='=HYPERLINK("http://evil","klik")')])

    cell = load_workbook(path).active.cell(row=2, column=2)
    assert cell.data_type == "s"
    assert cell.value.startswith("=")


def test_sheets_header_mismatch_is_not_overwritten(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    worksheet = _FakeWorksheet()
    worksheet.header = ["Header Lama"]
    monkeypatch.setattr(gspread, "service_account", lambda filename: _FakeClient(worksheet))

    GoogleSheetsWriter(GoogleSheetsConfig(spreadsheet_id="abc")).ensure_header()

    assert worksheet.header == ["Header Lama"]


class _BrokenWorksheet(_FakeWorksheet):
    def append_rows(self, rows: list[list[str | int]], value_input_option: str = "") -> None:
        raise GSpreadException("quota exceeded")


def test_sheets_append_error_is_wrapped(monkeypatch: pytest.MonkeyPatch) -> None:
    worksheet = _BrokenWorksheet()
    monkeypatch.setattr(gspread, "service_account", lambda filename: _FakeClient(worksheet))

    writer = GoogleSheetsWriter(GoogleSheetsConfig(spreadsheet_id="abc"))
    with pytest.raises(WriterError, match="Gagal menulis"):
        writer.append([_job()])


def test_excel_writer_rejects_corrupt_file(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    path.write_text("not a zip", encoding="utf-8")
    with pytest.raises(WriterError, match="tidak bisa dibaca"):
        ExcelWriter(path).ensure_header()


def test_build_writer_google_sheets_branch(monkeypatch: pytest.MonkeyPatch) -> None:
    class _FakeSheetsWriter:
        def __init__(self, cfg: GoogleSheetsConfig) -> None:
            self.cfg = cfg

        def ensure_header(self) -> None:
            return None

    monkeypatch.setattr(writers_pkg, "GoogleSheetsWriter", _FakeSheetsWriter)
    cfg = OutputConfig(type="google_sheets", google_sheets=GoogleSheetsConfig(spreadsheet_id="abc"))
    assert isinstance(build_writer(cfg), _FakeSheetsWriter)


def test_build_writer_missing_excel_section() -> None:
    cfg = OutputConfig.model_construct(type="excel", excel=None, google_sheets=None)
    with pytest.raises(WriterError, match=r"output\.excel"):
        build_writer(cfg)


def test_build_writer_missing_sheets_section() -> None:
    cfg = OutputConfig.model_construct(type="google_sheets", excel=None, google_sheets=None)
    with pytest.raises(WriterError, match=r"output\.google_sheets"):
        build_writer(cfg)


def test_null_writer_is_noop() -> None:
    writer = NullWriter()
    writer.ensure_header()
    writer.append([_job()])
